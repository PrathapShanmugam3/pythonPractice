"""Excel progress tracker. Re-running keeps the learner's own progress (matched by topic title)."""
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HEADERS = ["S.No", "Topic", "Level", "Notes PDF", "Program Files", "Status", "Date Started",
           "Date Completed", "Confidence (1-5)", "My Notes"]
LEVELS = ["Beginner", "Intermediate", "Advanced", "Master"]
USER_COLS = (6, 7, 8, 9, 10)  # Status .. My Notes (1-based)


def _existing_progress(path):
    if not path.exists():
        return {}
    ws = load_workbook(path)["Tracker"]
    return {row[1]: row[5:10] for row in ws.iter_rows(min_row=2, values_only=True) if row[1]}


def write_tracker(path, subject, rows):
    """rows: list of dict(topic, level, notes_pdf, programs[str])"""
    old = _existing_progress(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Tracker"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="DD0031")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, r in enumerate(rows, 1):
        kept = old.get(r["topic"])
        progress = list(kept) if kept else ["Not Started", None, None, None, None]
        ws.append([i, r["topic"], r["level"], r["notes_pdf"],
                   ", ".join(r["programs"]) if r["programs"] else "-"] + progress)
        row = i + 1
        ws.cell(row, 4).hyperlink = f"Notes/{r['notes_pdf']}"
        ws.cell(row, 4).font = Font(color="0563C1", underline="single")
        for col in (7, 8):
            ws.cell(row, col).number_format = "dd-mmm-yyyy"
    last = len(rows) + 1
    for col, w in zip("ABCDEFGHIJ", (6, 38, 14, 40, 44, 14, 14, 15, 12, 40)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, max_row=last):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:J{last}"
    status = DataValidation(type="list", formula1='"Not Started,In Progress,Done"', allow_blank=False)
    conf = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True,
                          errorTitle="Confidence", error="Enter a whole number from 1 to 5")
    dates = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                           errorTitle="Date", error="Enter a valid date")
    for dv, rng in ((status, f"F2:F{last}"), (conf, f"I2:I{last}"), (dates, f"G2:H{last}")):
        ws.add_data_validation(dv)
        dv.add(rng)
    for text, color in (("Done", "C6EFCE"), ("In Progress", "FFEB9C"), ("Not Started", "F2F2F2")):
        ws.conditional_formatting.add(f"F2:F{last}", CellIsRule(operator="equal", formula=[f'"{text}"'],
                                                                fill=PatternFill("solid", bgColor=color)))

    s = wb.create_sheet("Summary")
    s.append([f"{subject} - Progress Summary"])
    s["A1"].font = Font(bold=True, size=14, color="DD0031")
    s.append([])
    s.append(["Level", "Total Topics", "Not Started", "In Progress", "Done", "% Completed"])
    for c in s[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="DD0031")
        c.alignment = Alignment(horizontal="center")
    rng_lvl, rng_st = f"Tracker!$C$2:$C${last}", f"Tracker!$F$2:$F${last}"
    for i, lvl in enumerate(LEVELS):
        r = 4 + i
        s.append([lvl, f'=COUNTIF({rng_lvl},A{r})',
                  f'=COUNTIFS({rng_lvl},A{r},{rng_st},"Not Started")',
                  f'=COUNTIFS({rng_lvl},A{r},{rng_st},"In Progress")',
                  f'=COUNTIFS({rng_lvl},A{r},{rng_st},"Done")',
                  f"=IF(B{r}=0,0,E{r}/B{r})"])
    r = 4 + len(LEVELS)
    s.append(["OVERALL"] + [f"=SUM({c}4:{c}{r-1})" for c in "BCDE"] + [f"=IF(B{r}=0,0,E{r}/B{r})"])
    for c in s[r]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="F4F4F6")
    for rr in range(4, r + 1):
        s.cell(rr, 6).number_format = "0%"
    s.append([])
    s.append(["Average confidence (topics rated)", f'=IFERROR(AVERAGE(Tracker!$I$2:$I${last}),"-")'])
    s.cell(r + 2, 2).number_format = "0.0"
    for col, w in zip("ABCDEF", (34, 14, 14, 14, 10, 14)):
        s.column_dimensions[col].width = w
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
