"""Builds all PDFs + the Excel workbook for AIML_Learning from tools/content/*.py.

Run:  python3 tools/build.py   (from AIML_Learning/)
Needs: pip install reportlab openpyxl
"""
import importlib, os, pkgutil, sys, urllib.parse
from xml.sax.saxutils import escape
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer,
                                Preformatted, PageBreak, KeepTogether, Table, TableStyle)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(HERE, "content"))
F = "/usr/share/fonts/truetype/dejavu/"
pdfmetrics.registerFont(TTFont("S", F + "DejaVuSans.ttf"))
pdfmetrics.registerFont(TTFont("SB", F + "DejaVuSans-Bold.ttf"))
pdfmetrics.registerFont(TTFont("M", F + "DejaVuSansMono.ttf"))
pdfmetrics.registerFontFamily("S", normal="S", bold="SB", italic="S", boldItalic="SB")

def yt(q):
    return "https://www.youtube.com/results?search_query=" + urllib.parse.quote_plus(q)

def P(name, **kw):
    return ParagraphStyle(name, fontName=kw.pop("fontName", "S"), **kw)
H0 = P("h0", fontName="SB", fontSize=22, leading=28, textColor=colors.HexColor("#0b3d91"), spaceAfter=6)
H1 = P("h1", fontName="SB", fontSize=14, leading=18, textColor=colors.white, backColor=colors.HexColor("#0b3d91"),
       borderPadding=(4, 4, 4, 4), spaceBefore=12, spaceAfter=8)
LBL = P("lbl", fontName="SB", fontSize=9, leading=12, textColor=colors.HexColor("#0b3d91"), spaceBefore=4)
BODY = P("b", fontSize=9, leading=13)
TAN = P("t", fontSize=9, leading=13, backColor=colors.HexColor("#fff4e0"), borderPadding=3)
CODE = P("c", fontName="M", fontSize=7.5, leading=9.5, backColor=colors.HexColor("#f2f4f7"), borderPadding=3)
LINK = P("l", fontSize=8, leading=11, textColor=colors.HexColor("#1a0dab"))

def footer(c, d):
    c.saveState(); c.setFont("S", 7.5); c.setFillColor(colors.grey)
    c.drawString(15*mm, 8*mm, "AIML Learning - Zero to Master | Tamil+English (Tanglish)")
    c.drawRightString(A4[0]-15*mm, 8*mm, "Page %d" % d.page); c.restoreState()

def make_pdf(path, title, desc, topics, extra=()):
    doc = BaseDocTemplate(path, pagesize=A4, leftMargin=15*mm, rightMargin=15*mm, topMargin=15*mm,
                          bottomMargin=15*mm, title=title, author="AIML Learning")
    doc.addPageTemplates([PageTemplate(id="p", frames=[Frame(doc.leftMargin, doc.bottomMargin, doc.width,
                                                             doc.height, id="f")], onPage=footer)])
    s = [Paragraph(escape(title), H0), Paragraph(escape(desc), BODY), Spacer(1, 6)]
    rows = [["#", "Topic", "Level"]] + [[str(i), t["name"], t["level"]] for i, t in enumerate(topics, 1)]
    tb = Table([[Paragraph(escape(c), BODY) for c in r] for r in rows], colWidths=[10*mm, 130*mm, 35*mm], repeatRows=1)
    tb.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dbe5f7")),
                            ("GRID", (0, 0), (-1, -1), .25, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    s += [Paragraph("Contents", LBL), tb, PageBreak()]
    for i, t in enumerate(topics, 1):
        blk = [Paragraph("%d. %s  [%s]" % (i, escape(t["name"]), t["level"]), H1),
               Paragraph("Explanation (English)", LBL), Paragraph(escape(t["en"]), BODY),
               Paragraph("Tanglish (Tamil + English) Vilakkam", LBL), Paragraph(escape(t["ta"]), TAN)]
        if t["code"]:
            blk += [Paragraph("Example / Program", LBL), Preformatted(t["code"].strip("\n"), CODE)]
        blk += [Paragraph("Learn more", LBL),
                Paragraph('YouTube search: <a href="%s" color="#1a0dab">%s</a>' % (escape(yt(t["yt"])), escape(t["yt"])), LINK)]
        for lab, url in t.get("refs", []):
            blk.append(Paragraph('%s: <a href="%s" color="#1a0dab">%s</a>' % (escape(lab), escape(url), escape(url)), LINK))
        s.append(KeepTogether(blk[:4])); s += blk[4:]
    s += list(extra)
    doc.build(s)

def topic(name, level, en, ta, code="", yt_q=None, refs=()):
    return dict(name=name, level=level, en=en, ta=ta, code=code, yt=yt_q or name + " tutorial", refs=list(refs))

def main():
    import content_index
    mods = content_index.load()
    os.makedirs(os.path.join(ROOT, "pdfs"), exist_ok=True)
    wb = Workbook(); hdr_fill = PatternFill("solid", fgColor="0B3D91"); hdr_font = Font(bold=True, color="FFFFFF")
    def sheet(ws, headers, rows, widths):
        ws.append(headers)
        for c in ws[1]: c.fill, c.font = hdr_fill, hdr_font
        for r in rows: ws.append(r)
        for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    ws = wb.active; ws.title = "All_Topics"
    rows = []; summary = []
    for m in mods:
        fn = "%s.pdf" % m.FILE
        make_pdf(os.path.join(ROOT, "pdfs", fn), m.TITLE, m.DESC, m.TOPICS, getattr(m, "EXTRA", ()))
        summary.append([m.NUM, m.TITLE, len(m.TOPICS), "pdfs/" + fn])
        for i, t in enumerate(m.TOPICS, 1):
            code1 = (t["code"].strip().splitlines() or [""])[0] if t["code"] else ""
            rows.append([m.NUM, m.TITLE, i, t["name"], t["level"], t["en"], t["ta"], t["code"].strip(), yt(t["yt"]),
                         "; ".join(u for _, u in t["refs"]), "Not Started", ""])
    sheet(ws, ["Module", "Module Title", "#", "Topic", "Level", "Explanation (English)", "Tanglish Explanation",
               "Example / Program", "YouTube Search Link", "Other References", "Status", "My Notes"], rows,
          [8, 24, 5, 30, 13, 60, 60, 60, 40, 40, 14, 30])
    dv = DataValidation(type="list", formula1='"Not Started,In Progress,Done,Revise"', allow_blank=True)
    ws.add_data_validation(dv); dv.add("K2:K%d" % (len(rows) + 1))
    ws2 = wb.create_sheet("Modules_Summary", 0)
    sheet(ws2, ["Module", "Title", "Topics", "PDF File"], summary, [8, 50, 8, 45])
    import resources
    for name, headers, data, widths in resources.SHEETS:
        sheet(wb.create_sheet(name), headers, data, widths)
    wb.save(os.path.join(ROOT, "AIML_Learning_Tracker.xlsx"))
    print("modules:", len(mods), "topics:", len(rows))

if __name__ == "__main__":
    main()
